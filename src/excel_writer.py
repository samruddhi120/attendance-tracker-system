"""
Exports attendance from SQLite into an Excel workbook - one sheet per date
inside a single master workbook, matching the schema from the project plan
(Roll No, Name, Date, Time, Status, Source).

SQLite is the source of truth here; this script only ever reads from the DB
and regenerates the sheet, so the DB and the Excel report can never drift
apart or need manual reconciliation.
"""

import os

from openpyxl import Workbook, load_workbook

from config import MASTER_WORKBOOK
from db import get_attendance_for_date, get_all_students

COLUMNS = ["Roll No", "Name", "Date", "Time", "Status", "Source"]


def export_day(date):
    """
    Writes one row per ENROLLED STUDENT, not just per attendance record -
    students with no matching attendance row get an explicit "Absent" row
    so the sheet is always the full class roster, not just who showed up.
    """
    os.makedirs(os.path.dirname(MASTER_WORKBOOK), exist_ok=True)

    if os.path.exists(MASTER_WORKBOOK):
        workbook = load_workbook(MASTER_WORKBOOK)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)  # drop the default empty sheet

    sheet_name = date  # e.g. "2026-08-30"
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]  # regenerate fresh each run - no stale duplicate rows
    sheet = workbook.create_sheet(sheet_name)
    sheet.append(COLUMNS)

    present_rows = {row[0]: row for row in get_attendance_for_date(date)}  # keyed by roll_no
    all_students = get_all_students()

    present_count = 0
    for roll_no, name in all_students:
        if roll_no in present_rows:
            _roll_no, _name, date_, time_, status, source = present_rows[roll_no]
            sheet.append([roll_no, name, date_, time_, status, source])
            present_count += 1
        else:
            sheet.append([roll_no, name, date, "", "Absent", "-"])

    workbook.save(MASTER_WORKBOOK)
    print(
        f"Exported {len(all_students)} student(s) for {date} to {MASTER_WORKBOOK} "
        f"(sheet '{sheet_name}') - {present_count} present, {len(all_students) - present_count} absent."
    )


if __name__ == "__main__":
    from datetime import datetime

    export_day(datetime.now().strftime("%Y-%m-%d"))
